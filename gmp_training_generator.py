"""
GMP SOP 교육자료 및 퀴즈 자동 생성 모듈 (분할 생성 버전)
React 대시보드 백엔드에서 사용할 수 있는 독립 모듈
OpenAI GPT-4 사용 - 토큰 제한 극복을 위한 분할 전략
"""

import os
import json
import time
from typing import Dict, List, Optional, Callable
from dataclasses import dataclass, asdict
from docx import Document
from openai import OpenAI
from dotenv import load_dotenv

load_dotenv()


@dataclass
class TrainingMaterial:
    """교육자료 데이터 구조"""
    title: str
    summary: str
    key_points: List[str]
    detailed_content: str
    references: List[str]


@dataclass
class QuizQuestion:
    """퀴즈 문제 데이터 구조"""
    question: str
    options: List[str]
    correct_answer: int
    explanation: str
    reference_section: str


@dataclass
class GeneratedOutput:
    """생성된 결과물 전체 구조"""
    training_material: TrainingMaterial
    quiz_questions: List[QuizQuestion]
    metadata: Dict


class GMPTrainingGenerator:
    """GMP SOP 교육자료 및 퀴즈 생성기 (분할 생성 버전)"""
    
    def __init__(self, api_key: Optional[str] = None, model: str = "gpt-4-turbo-preview"):
        if api_key is None:
            api_key = os.getenv("OPENAI_API_KEY")
            if not api_key:
                raise ValueError("OPENAI_API_KEY가 설정되지 않았습니다.")
        
        self.client = OpenAI(api_key=api_key)
        self.model = model
    
    def extract_text_from_docx(self, file_path: str) -> str:
        """Word 문서에서 텍스트 추출"""
        doc = Document(file_path)
        paragraphs = (para.text for para in doc.paragraphs if para.text.strip())
        
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        table_texts.append(cell.text)
        
        all_text = list(paragraphs) + table_texts
        return "\n".join(all_text)
    
    def generate_training_material(
        self, 
        sop_content: str, 
        guideline_content: str,
        additional_context: Optional[str] = None,
        max_retries: int = 2
    ) -> TrainingMaterial:
        """
        교육자료 생성 (2단계 분할 생성)
        1단계: 개요, 요약, 핵심포인트, 참조 생성
        2단계: 상세 내용만 별도로 생성
        """
        
        print("   [1단계] 개요 및 핵심 포인트 생성 중...")
        overview = self._generate_overview(sop_content, guideline_content, additional_context)
        
        print("   [2단계] 상세 내용 생성 중 (이 단계가 오래 걸립니다)...")
        detailed_content = self._generate_detailed_content(
            sop_content, 
            guideline_content, 
            overview,
            additional_context
        )
        
        return TrainingMaterial(
            title=overview['title'],
            summary=overview['summary'],
            key_points=overview['key_points'],
            detailed_content=detailed_content,
            references=overview['references']
        )
    
    def _generate_overview(
        self, 
        sop_content: str, 
        guideline_content: str,
        additional_context: Optional[str] = None
    ) -> Dict:
        """1단계: 개요, 요약, 핵심포인트만 생성"""
        
        additional_context_section = ""
        if additional_context:
            additional_context_section = f"\n# 추가 컨텍스트:\n{additional_context}\n"
        
        prompt = f"""당신은 GMP 전문가입니다. 
제공된 SOP와 가이드라인을 분석하여 개요를 작성해주세요.

# SOP 내용:
{sop_content[:2000]}

# 가이드라인 내용:
{guideline_content[:2000]}{additional_context_section}

아래 JSON 형식으로 작성하세요:
{{
  "title": "교육자료 제목 (명확하고 구체적으로)",
  "summary": "3-4문장의 상세한 요약 (각 문장은 완전한 정보 포함)",
  "key_points": [
    "핵심 포인트 1 (구체적이고 실무 중심, 한 문장)",
    "핵심 포인트 2",
    "..."
  ],
  "references": ["참조 섹션 1", "참조 섹션 2", ...]
}}

**요구사항:**
- key_points: 7-10개, 각 포인트는 SOP의 주요 섹션을 대표
- summary: 이 SOP의 목적, 범위, 주요 규제 근거를 간결하게 설명
- references: SOP와 가이드라인의 모든 주요 섹션 나열

반드시 유효한 JSON으로만 응답하세요."""

        response = self.client.chat.completions.create(
            model=self.model,
            messages=[
                {"role": "system", "content": "당신은 GMP 전문가입니다. 항상 유효한 JSON으로 응답합니다."},
                {"role": "user", "content": prompt}
            ],
            temperature=0.2,
            max_tokens=2000,
            response_format={"type": "json_object"}
        )
        
        response_text = response.choices[0].message.content
        data = json.loads(response_text)
        
        print(f"      ✓ 제목: {data['title']}")
        print(f"      ✓ 핵심 포인트: {len(data['key_points'])}개")
        
        return data
    
    def _generate_detailed_content(
        self,
        sop_content: str,
        guideline_content: str,
        overview: Dict,
        additional_context: Optional[str] = None
    ) -> str:
        """2단계: 상세 내용만 생성 (토큰을 상세 내용에만 집중)"""
        
        additional_context_section = ""
        if additional_context:
            additional_context_section = f"\n# 추가 컨텍스트:\n{additional_context}\n"
        
        # 핵심 포인트를 가이드로 제공
        key_points_str = "\n".join([f"- {kp}" for kp in overview['key_points']])
        
        prompt = f"""당신은 GMP 전문가입니다.
아래 SOP와 가이드라인을 바탕으로 실무자를 위한 상세한 교육 내용을 마크다운 형식으로 작성하세요.

# SOP 내용:
{sop_content[:3500]}

# 가이드라인 내용:
{guideline_content[:3500]}{additional_context_section}

# 다룰 핵심 주제:
{key_points_str}

다음 구조로 작성하되, **각 내용마다 반드시 출처를 명시**하세요:

## 1. 개요

SOP의 목적, 적용 범위, 관련 규정 근거를 자연스러운 문장으로 설명하세요.
(예: "본 SOP는 21 CFR 211.122에 근거하여...")

## 2. 주요 내용

위의 각 핵심 주제에 대해 다음 형식으로 작성:

### [주제명]

**SOP 근거**: (SOP Section X에서...)
- 요구사항 1
- 요구사항 2
- 담당 부서 및 역할

**가이드라인 근거**: (21 CFR 211.XXX 또는 해당 조항)
- 가이드라인의 원문 또는 주요 내용
- 이 규정이 요구하는 이유와 배경
- 준수하지 않을 경우의 위험

**실무 적용**:
1. [구체적 단계] - 어떻게 수행하는지 상세히
2. [다음 단계] - 필요한 문서나 양식
3. [마지막 단계] - 확인 및 기록 방법

**예시 또는 주의사항**: 
- 실제 업무에서 놓치기 쉬운 부분
- 흔한 실수와 예방법

각 섹션은 400-600자로 작성하고, 문장은 자연스럽게 연결하세요.

## 3. 체크리스트

주요 절차별 확인사항을 구체적으로 나열하세요.

## 4. 주의사항 및 흔한 실수

실무에서 자주 발생하는 오류 5가지 이상을 구체적 상황과 함께 설명하세요.

**작성 원칙:**
- 모든 내용은 자연스러운 문장으로 작성 (나열식 금지)
- 출처를 명확히 (예: "SOP Section 3에 따르면...", "21 CFR 211.125는...")
- SOP와 가이드라인을 직접 연결하여 설명
- 생략 표현("..." 등) 사용 금지

마크다운으로만 응답하세요."""

        response = self.client.chat.completions.create(
            model=self.model,
            messages=[
                {
                    "role": "system", 
                    "content": "당신은 GMP 전문가입니다. 실무자가 읽기 쉽도록 자연스러운 문장으로 교육 자료를 작성합니다. 모든 내용에 출처(SOP 섹션 번호, 가이드라인 조항)를 명시합니다. 마크다운 형식으로만 응답합니다."
                },
                {"role": "user", "content": prompt}
            ],
            temperature=0.4,
            max_tokens=4000,
        )
        
        detailed_content = response.choices[0].message.content.strip()
        
        print(f"      ✓ 상세 내용 생성 완료: {len(detailed_content):,}자")
        
        return detailed_content
    
    def generate_quiz(
        self, 
        sop_content: str, 
        guideline_content: str,
        num_questions: int = 10,
        difficulty: str = "mixed",
        max_retries: int = 2
    ) -> List[QuizQuestion]:
        """퀴즈 생성"""
        
        if difficulty == "mixed":
            easy_count = max(1, int(num_questions * 0.3))
            hard_count = max(1, int(num_questions * 0.2))
            medium_count = num_questions - easy_count - hard_count
            difficulty_instruction = f"반드시 easy {easy_count}문제, medium {medium_count}문제, hard {hard_count}문제를 정확히 생성하세요."
        else:
            difficulty_instruction = f"모든 문제를 {difficulty} 난이도로 생성하세요."
        
        prompt = f"""당신은 GMP 교육 전문가입니다.
제공된 SOP와 가이드라인을 바탕으로 평가용 퀴즈를 작성해주세요.

# SOP 내용:
{sop_content[:2000]}

# 가이드라인 내용:
{guideline_content[:2000]}

{num_questions}개의 객관식 문제를 생성해주세요.
{difficulty_instruction}

JSON 형식:
{{
  "questions": [
    {{
      "question": "문제 내용",
      "options": ["선택지1", "선택지2", "선택지3", "선택지4", "선택지5"],
      "correct_answer": 0,
      "explanation": "정답 해설 (150자 이상)",
      "reference_section": "참조 섹션"
    }}
  ]
}}

**필수 요구사항:**
1. 반드시 5개 선택지
2. 정답은 0-4 인덱스
3. 난이도 구분:
   - easy: 정의, 기본 개념
   - medium: 절차 이해, 적용
   - hard: 복합 상황 판단
4. 오답은 그럴듯하게 (완전히 틀린 답 금지)
5. 해설은 상세하게 (왜 정답인지, 다른 답은 왜 틀렸는지)

반드시 유효한 JSON으로만 응답하세요."""

        for attempt in range(max_retries):
            try:
                print(f"   시도 {attempt + 1}/{max_retries}...")
                
                response = self.client.chat.completions.create(
                    model=self.model,
                    messages=[
                        {"role": "system", "content": "당신은 GMP 전문가입니다. 유효한 JSON으로만 응답합니다."},
                        {"role": "user", "content": prompt}
                    ],
                    temperature=0.4,
                    max_tokens=3500,
                    response_format={"type": "json_object"}
                )
                
                response_text = response.choices[0].message.content
                data = json.loads(response_text)
                
                if self._validate_quiz(data, num_questions):
                    print(f"   ✓ 퀴즈 생성 완료!")
                    return [QuizQuestion(**q) for q in data["questions"]]
                else:
                    if attempt < max_retries - 1:
                        time.sleep(2)
                        continue
                    
            except Exception as e:
                print(f"   ⚠️  오류: {e}")
                if attempt < max_retries - 1:
                    time.sleep(2)
                    continue
                else:
                    raise
        
        raise Exception(f"퀴즈 생성 실패")
    
    def _validate_quiz(self, data: Dict, expected_count: int) -> bool:
        """퀴즈 품질 검증"""
        try:
            if 'questions' not in data:
                return False
            
            questions = data['questions']
            if len(questions) != expected_count:
                print(f"      문제 수 불일치: {len(questions)}개")
                return False
            
            for i, q in enumerate(questions, 1):
                required_fields = ['question', 'options', 'correct_answer', 'explanation', 'reference_section']
                if not all(field in q for field in required_fields):
                    print(f"      문제 {i}: 필수 필드 누락")
                    return False
                
                if len(q['options']) != 5:
                    print(f"      문제 {i}: 선택지 개수 오류")
                    return False
                
                if not isinstance(q['correct_answer'], int) or q['correct_answer'] < 0 or q['correct_answer'] > 4:
                    print(f"      문제 {i}: 정답 인덱스 오류")
                    return False
            
            return True
            
        except Exception as e:
            print(f"      검증 오류: {e}")
            return False
    
    def generate_from_files(
        self,
        sop_file_path: str,
        guideline_file_path: str,
        num_quiz_questions: int = 10,
        quiz_difficulty: str = "mixed",
        additional_context: Optional[str] = None,
        progress_callback: Optional[Callable[[str], None]] = None
    ) -> GeneratedOutput:
        """파일에서 직접 교육자료와 퀴즈 생성"""
        
        def log(message: str):
            print(message)
            if progress_callback:
                progress_callback(message)
        
        log("📄 SOP 파일 읽는 중...")
        sop_content = self.extract_text_from_docx(sop_file_path)
        
        log("📄 가이드라인 파일 읽는 중...")
        guideline_content = self.extract_text_from_docx(guideline_file_path)
        
        log("✍️  교육자료 생성 중 (2단계 분할 생성, 2-3분 소요)...")
        training_material = self.generate_training_material(
            sop_content, 
            guideline_content,
            additional_context
        )
        
        log(f"❓ 퀴즈 문제 {num_quiz_questions}개 생성 중...")
        quiz_questions = self.generate_quiz(
            sop_content,
            guideline_content,
            num_quiz_questions,
            quiz_difficulty
        )
        
        metadata = {
            "sop_file": os.path.basename(sop_file_path),
            "guideline_file": os.path.basename(guideline_file_path),
            "num_questions": len(quiz_questions),
            "difficulty": quiz_difficulty,
            "generation_timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
            "training_content_length": len(training_material.detailed_content)
        }
        
        log("✅ 생성 완료!")
        
        return GeneratedOutput(
            training_material=training_material,
            quiz_questions=quiz_questions,
            metadata=metadata
        )
    
    def export_to_separate_json(self, output: GeneratedOutput, base_filename: str = "gmp_training"):
        """교육자료와 퀴즈를 별도 JSON 파일로 내보내기"""
        training_file = f"{base_filename}_교육자료.json"
        training_data = {
            "training_material": asdict(output.training_material),
            "metadata": output.metadata
        }
        with open(training_file, 'w', encoding='utf-8') as f:
            json.dump(training_data, f, ensure_ascii=False, indent=2)
        print(f"💾 교육자료 저장 완료: {training_file}")
        
        quiz_file = f"{base_filename}_퀴즈.json"
        quiz_data = {
            "quiz_questions": [asdict(q) for q in output.quiz_questions],
            "metadata": output.metadata
        }
        with open(quiz_file, 'w', encoding='utf-8') as f:
            json.dump(quiz_data, f, ensure_ascii=False, indent=2)
        print(f"💾 퀴즈 저장 완료: {quiz_file}")
        
        return training_file, quiz_file
    
    def to_dict(self, output: GeneratedOutput) -> Dict:
        """딕셔너리로 변환"""
        return asdict(output)